"""Admin endpoints for event approval and management"""

from fastapi import APIRouter, Depends, HTTPException, Query, Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_, func, or_
from pydantic import BaseModel
from datetime import datetime, timezone

from app.core.database import get_db
from app.core.security import get_current_user
from app.core.audit import log_action
from app.models.user import User
from app.models.event import Event
from app.models.guest import Guest
from app.services.notify import send_notification

router = APIRouter()


class EventApprovalRequest(BaseModel):
    event_id: int
    approved: bool
    reason: str | None = None


class EventReviewResponse(BaseModel):
    id: int
    title: str
    host_name: str
    event_type: str
    review_status: str
    flagged_keywords: list[str] | None
    organizer_email: str
    created_at: datetime

    class Config:
        from_attributes = True


async def check_admin(user: User = Depends(get_current_user)) -> User:
    """Verify user is admin"""
    if user.role not in ["admin", "super_admin"]:
        raise HTTPException(status_code=403, detail="Admin access required")
    return user


@router.get("/admin/events/pending")
async def list_pending_events(
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
):
    """List all events pending admin review"""
    offset = (page - 1) * limit

    result = await db.execute(
        select(Event)
        .where(Event.review_status == "pending_review")
        .order_by(Event.created_at.desc())
        .limit(limit)
        .offset(offset)
    )
    events = result.scalars().all()

    total_result = await db.execute(
        select(func.count(Event.id)).where(Event.review_status == "pending_review")
    )
    total = total_result.scalar() or 0
    total_pages = (total + limit - 1) // limit

    return {
        "events": [
            {
                "id": e.id,
                "title": e.title,
                "host_name": e.host_name,
                "event_type": e.event_type,
                "event_date": e.event_date,
                "event_time": e.event_time,
                "venue": e.venue,
                "ticket_price": e.ticket_price,
                "status": e.review_status,
                "created_at": e.created_at,
            }
            for e in events
        ],
        "total": total,
        "total_pages": total_pages,
        "page": page,
        "limit": limit,
    }


@router.get("/admin/events/flagged")
async def list_flagged_events(
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
):
    """List all flagged events for manual review"""
    offset = (page - 1) * limit

    result = await db.execute(
        select(Event)
        .where(Event.review_status == "flagged")
        .order_by(Event.created_at.desc())
        .limit(limit)
        .offset(offset)
    )
    events = result.scalars().all()

    total_result = await db.execute(
        select(func.count(Event.id)).where(Event.review_status == "flagged")
    )
    total = total_result.scalar() or 0
    total_pages = (total + limit - 1) // limit

    return {
        "events": [
            {
                "id": e.id,
                "title": e.title,
                "host_name": e.host_name,
                "event_date": e.event_date,
                "event_time": e.event_time,
                "venue": e.venue,
                "ticket_price": e.ticket_price,
                "status": e.review_status,
                "created_at": e.created_at,
            }
            for e in events
        ],
        "total": total,
        "total_pages": total_pages,
        "page": page,
        "limit": limit,
    }


@router.post("/admin/events/{event_id}/approve")
async def approve_event(
    event_id: int,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Approve an event for public listing"""
    result = await db.execute(select(Event).where(Event.id == event_id))
    event = result.scalar_one_or_none()

    if not event:
        raise HTTPException(status_code=404, detail="Event not found")

    if event.review_status not in ["pending_review", "flagged"]:
        raise HTTPException(status_code=400, detail="Event is not pending review")

    # Update event status
    event.status = "published"
    event.is_public = True
    event.review_status = "approved"
    event.updated_at = datetime.now(timezone.utc)

    db.add(event)
    await db.commit()
    await db.refresh(event)

    # Notify organizer
    organizer_result = await db.execute(
        select(User).where(User.id == event.organizer_id)
    )
    organizer = organizer_result.scalar_one_or_none()

    if organizer:
        await send_notification(
            db=db,
            user_id=organizer.id,
            type="event_approved",
            title="Your Event Was Approved! 🎉",
            message=f"Your event '{event.title}' has been approved and is now live on Discover Events.",
            data={"event_id": event.id},
        )

    await log_action(
        db=db, user_id=admin.id, action="event_approved",
        resource_type="event", resource_id=event.id,
        description=f"Event '{event.title}' approved and published",
    )

    return {
        "status": "approved",
        "event_id": event.id,
        "message": f"Event '{event.title}' approved and published",
    }


class RejectRequest(BaseModel):
    reason: str = "No reason provided"


@router.post("/admin/events/{event_id}/reject")
async def reject_event(
    event_id: int,
    req: RejectRequest,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Reject an event from public listing"""
    result = await db.execute(select(Event).where(Event.id == event_id))
    event = result.scalar_one_or_none()

    if not event:
        raise HTTPException(status_code=404, detail="Event not found")

    if event.review_status not in ["pending_review", "flagged"]:
        raise HTTPException(status_code=400, detail="Event is not pending review")

    # Update event status
    event.status = "rejected"
    event.is_public = False
    event.review_status = "rejected"
    event.review_note = req.reason
    event.updated_at = datetime.now(timezone.utc)

    db.add(event)
    await db.commit()
    await db.refresh(event)

    # Notify organizer
    organizer_result = await db.execute(
        select(User).where(User.id == event.organizer_id)
    )
    organizer = organizer_result.scalar_one_or_none()

    if organizer:
        await send_notification(
            db=db,
            user_id=organizer.id,
            type="event_rejected",
            title="Event Submission Requires Revision",
            message=f"Your event '{event.title}' was not approved. Reason: {req.reason}",
            data={"event_id": event.id, "reason": req.reason},
        )

    await log_action(
        db=db, user_id=admin.id, action="event_rejected",
        resource_type="event", resource_id=event.id,
        description=f"Event '{event.title}' rejected. Reason: {reason}",
    )

    return {
        "status": "rejected",
        "event_id": event.id,
        "reason": reason,
        "message": f"Event '{event.title}' rejected",
    }


@router.delete("/admin/events/{event_id}")
async def delete_event(
    event_id: int,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Delete an event (admin only)"""
    result = await db.execute(select(Event).where(Event.id == event_id))
    event = result.scalar_one_or_none()

    if not event:
        raise HTTPException(status_code=404, detail="Event not found")

    # Delete the event
    db.delete(event)
    await db.commit()

    await log_action(
        db=db, user_id=admin.id, action="event_deleted",
        resource_type="event", resource_id=event_id,
        description=f"Event '{event.title}' deleted by admin",
    )

    return {
        "status": "deleted",
        "event_id": event_id,
        "message": "Event deleted successfully",
    }


class EventControlsUpdate(BaseModel):
    guest_count_range: str | None = None
    registration_open: bool | None = None
    registration_close_at: str | None = None
    rsvp_open: bool | None = None
    is_public: bool | None = None
    status: str | None = None
    ticket_price: int | None = None
    tickets_available: int | None = None
    spotlight: bool | None = None
    title: str | None = None
    venue: str | None = None
    city: str | None = None
    event_date: str | None = None
    event_time: str | None = None
    description: str | None = None


def _controls_dict(event: Event) -> dict:
    return {
        "id": event.id,
        "title": event.title,
        "event_date": event.event_date,
        "event_time": event.event_time,
        "venue": event.venue,
        "city": getattr(event, "city", None),
        "description": getattr(event, "description", None),
        "organizer_id": event.organizer_id,
        "guest_count_range": event.guest_count_range,
        "registration_open": event.registration_open,
        "registration_close_at": event.registration_close_at,
        "rsvp_open": getattr(event, "rsvp_open", True),
        "is_public": event.is_public,
        "status": event.status,
        "ticket_price": event.ticket_price,
        "tickets_available": event.tickets_available,
        "spotlight": getattr(event, "spotlight", False),
        "invite_subject": event.invite_subject,
        "invite_body": event.invite_body,
    }


@router.get("/admin/events/search")
async def admin_search_events(
    q: str = Query("", description="Event id or title fragment"),
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Find any event by id or title for the admin controls panel."""
    stmt = select(Event).order_by(Event.id.desc()).limit(20)
    if q.strip().isdigit():
        stmt = select(Event).where(Event.id == int(q.strip()))
    elif q.strip():
        stmt = select(Event).where(Event.title.ilike(f"%{q.strip()}%")).order_by(Event.id.desc()).limit(20)
    result = await db.execute(stmt)
    return {"events": [_controls_dict(e) for e in result.scalars().all()]}


@router.patch("/admin/events/{event_id}/controls")
async def admin_update_event_controls(
    event_id: int,
    req: EventControlsUpdate,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Admin override: guest cap, registration open/close, visibility, status."""
    result = await db.execute(select(Event).where(Event.id == event_id))
    event = result.scalar_one_or_none()
    if not event:
        raise HTTPException(status_code=404, detail="Event not found")

    changes = []
    if req.guest_count_range is not None:
        import re as _re

        val = req.guest_count_range.strip()
        if val.lower() != "unlimited" and not _re.search(r"\d", val):
            raise HTTPException(status_code=400, detail="Guest cap must contain a number or be 'Unlimited'")
        changes.append(f"guest cap {event.guest_count_range} -> {val}")
        event.guest_count_range = val if val.lower() != "unlimited" else "Unlimited"
    if req.registration_open is not None:
        changes.append(f"registration {'opened' if req.registration_open else 'closed'}")
        event.registration_open = req.registration_open
    if req.registration_close_at is not None:
        from datetime import datetime

        raw = req.registration_close_at.strip()
        if raw == "":
            event.registration_close_at = None
            changes.append("scheduled registration close cleared")
        else:
            try:
                close_at = datetime.fromisoformat(raw.replace("Z", "+00:00"))
            except ValueError:
                raise HTTPException(status_code=400, detail="Invalid datetime (use ISO format)")
            event.registration_close_at = close_at
            changes.append(f"registration auto-closes at {close_at.isoformat()}")
    if req.rsvp_open is not None:
        changes.append(f"RSVP {'opened' if req.rsvp_open else 'closed'}")
        event.rsvp_open = req.rsvp_open
    if req.ticket_price is not None:
        if req.ticket_price < 0:
            raise HTTPException(status_code=400, detail="Ticket price cannot be negative")
        changes.append(f"ticket price {event.ticket_price} -> {req.ticket_price}")
        event.ticket_price = req.ticket_price
    if req.tickets_available is not None:
        if req.tickets_available < 0:
            raise HTTPException(status_code=400, detail="Tickets available cannot be negative")
        changes.append(f"tickets available {event.tickets_available} -> {req.tickets_available}")
        event.tickets_available = req.tickets_available
    if req.spotlight is not None:
        if req.spotlight:
            # Only one spotlight at a time: clear the flag everywhere else.
            await db.execute(
                Event.__table__.update().where(Event.spotlight == True).values(spotlight=False)  # noqa: E712
            )
        changes.append(f"spotlight {'on' if req.spotlight else 'off'}")
        event.spotlight = req.spotlight
    if req.title is not None and req.title.strip():
        changes.append(f"title -> {req.title.strip()[:60]}")
        event.title = req.title.strip()
    if req.venue is not None and req.venue.strip():
        changes.append("venue updated")
        event.venue = req.venue.strip()
    if req.city is not None:
        event.city = req.city.strip() or None
        changes.append("city updated")
    if req.event_date is not None and req.event_date.strip():
        from datetime import date as _date

        try:
            event.event_date = _date.fromisoformat(req.event_date.strip())
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid date (use YYYY-MM-DD)")
        changes.append(f"date -> {event.event_date}")
    if req.event_time is not None and req.event_time.strip():
        from datetime import time as _time

        try:
            parts = req.event_time.strip().split(":")
            event.event_time = _time(int(parts[0]), int(parts[1]))
        except (ValueError, IndexError):
            raise HTTPException(status_code=400, detail="Invalid time (use HH:MM)")
        changes.append(f"time -> {event.event_time}")
    if req.description is not None:
        event.description = req.description
        changes.append("description updated")
    if req.is_public is not None:
        changes.append(f"visibility -> {'public' if req.is_public else 'private'}")
        event.is_public = req.is_public
    if req.status is not None:
        if req.status not in ("draft", "published", "cancelled"):
            raise HTTPException(status_code=400, detail="Invalid status")
        changes.append(f"status {event.status} -> {req.status}")
        event.status = req.status

    if not changes:
        raise HTTPException(status_code=400, detail="No changes provided")

    await db.commit()
    await db.refresh(event)
    await log_action(
        db=db, user_id=admin.id, action="event_controls_updated",
        resource_type="event", resource_id=event_id,
        description="; ".join(changes),
    )
    return {"event": _controls_dict(event), "changes": changes}


class AdminGuestUpsert(BaseModel):
    name: str | None = None
    phone: str | None = None
    email: str | None = None
    category: str | None = None
    organization: str | None = None
    notes: str | None = None
    rsvp_status: str | None = None


async def _admin_get_event(db: AsyncSession, event_id: int) -> Event:
    result = await db.execute(select(Event).where(Event.id == event_id))
    event = result.scalar_one_or_none()
    if not event:
        raise HTTPException(status_code=404, detail="Event not found")
    return event


@router.get("/admin/events/{event_id}/guests")
async def admin_list_guests(
    event_id: int,
    search: str = Query(""),
    limit: int = Query(100, ge=1, le=1000),
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Admin guest list for any event + existing categories for suggestions."""
    await _admin_get_event(db, event_id)
    stmt = select(Guest).where(Guest.event_id == event_id, Guest.deleted_at == None)
    if search.strip():
        q = f"%{search.strip()}%"
        stmt = stmt.where(or_(Guest.name.ilike(q), Guest.email.ilike(q), Guest.phone.ilike(q)))
    stmt = stmt.order_by(Guest.id.desc()).limit(limit)
    guests = (await db.execute(stmt)).scalars().all()
    cat_rows = await db.execute(
        select(Guest.category, func.count())
        .where(Guest.event_id == event_id, Guest.deleted_at == None, Guest.category != None)
        .group_by(Guest.category)
    )
    return {
        "guests": [g.to_dict() for g in guests],
        "categories": [{"name": r[0], "count": r[1]} for r in cat_rows.all()],
    }


@router.post("/admin/events/{event_id}/guests")
async def admin_add_guest(
    event_id: int,
    req: AdminGuestUpsert,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Admin quick-add: same as gate quick-add but from the admin panel."""
    from app.api.guests import ensure_guest_capacity

    event = await _admin_get_event(db, event_id)
    name = (req.name or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Name is required")
    await ensure_guest_capacity(db, event, 1)
    custom = {"organization": req.organization.strip()} if (req.organization or "").strip() else None
    guest = Guest(
        event_id=event_id,
        name=name,
        phone=(req.phone or "").strip() or None,
        email=(req.email or "").strip() or None,
        category=(req.category or "").strip() or None,
        custom_data=custom,
        notes=(req.notes or "").strip() or None,
        rsvp_status=req.rsvp_status or "pending",
        invited_by="Admin",
    )
    db.add(guest)
    await db.commit()
    await db.refresh(guest)
    await log_action(db=db, user_id=admin.id, action="guest_added_by_admin",
                     resource_type="event", resource_id=event_id,
                     description=f"Added guest {guest.name} (id {guest.id})")
    return {"guest": guest.to_dict()}


@router.put("/admin/events/{event_id}/guests/{guest_id}")
async def admin_update_guest(
    event_id: int,
    guest_id: int,
    req: AdminGuestUpsert,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Admin edit: name, contacts, category, organization, notes, RSVP status."""
    await _admin_get_event(db, event_id)
    g = (await db.execute(select(Guest).where(Guest.id == guest_id, Guest.event_id == event_id))).scalar_one_or_none()
    if not g:
        raise HTTPException(status_code=404, detail="Guest not found")
    changes = []
    if req.name is not None and req.name.strip() and req.name.strip() != g.name:
        changes.append(f"name -> {req.name.strip()}")
        g.name = req.name.strip()
    if req.phone is not None:
        g.phone = req.phone.strip() or None
        changes.append("phone updated")
    if req.email is not None:
        g.email = req.email.strip() or None
        changes.append("email updated")
    if req.category is not None:
        g.category = req.category.strip() or None
        changes.append(f"category -> {g.category}")
    if req.organization is not None:
        custom = dict(g.custom_data or {})
        if req.organization.strip():
            custom["organization"] = req.organization.strip()
        else:
            custom.pop("organization", None)
        g.custom_data = custom
        changes.append("organization updated")
    if req.notes is not None:
        g.notes = req.notes.strip() or None
        changes.append("notes updated")
    if req.rsvp_status is not None:
        if req.rsvp_status not in ("pending", "accepted", "declined", "maybe"):
            raise HTTPException(status_code=400, detail="Invalid RSVP status")
        g.rsvp_status = req.rsvp_status
        changes.append(f"rsvp -> {req.rsvp_status}")
    await db.commit()
    await db.refresh(g)
    await log_action(db=db, user_id=admin.id, action="guest_updated_by_admin",
                     resource_type="event", resource_id=event_id,
                     description=f"Guest {g.id}: " + ("; ".join(changes) if changes else "no changes"))
    return {"guest": g.to_dict()}


@router.delete("/admin/events/{event_id}/guests/{guest_id}")
async def admin_delete_guest(
    event_id: int,
    guest_id: int,
    permanent: bool = Query(False),
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    await _admin_get_event(db, event_id)
    g = (await db.execute(select(Guest).where(Guest.id == guest_id, Guest.event_id == event_id))).scalar_one_or_none()
    if not g:
        raise HTTPException(status_code=404, detail="Guest not found")
    from sqlalchemy import text as _text

    if permanent:
        for table in ("invite_messages", "qr_codes", "checkins", "scan_attempts", "payments"):
            await db.execute(_text(f"DELETE FROM {table} WHERE guest_id = :gid"), {"gid": guest_id})
        await db.delete(g)
    else:
        g.deleted_at = func.now()
    await db.commit()
    await log_action(db=db, user_id=admin.id, action="guest_deleted_by_admin",
                     resource_type="event", resource_id=event_id,
                     description=f"Deleted guest {guest_id} ({'permanent' if permanent else 'soft'})")
    return {"message": "Guest deleted"}


@router.get("/admin/events/{event_id}/report-download")
async def admin_download_event_report(
    event_id: int,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """One-click post-event report: summary, guests, ticket sales and
    accreditation breakdown as a multi-sheet Excel file."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from fastapi.responses import StreamingResponse

    from app.models.checkin import CheckIn
    from app.models.ticket_purchase import TicketPurchase

    event = await _admin_get_event(db, event_id)
    guests = (await db.execute(
        select(Guest).where(Guest.event_id == event_id, Guest.deleted_at == None).order_by(Guest.id)
    )).scalars().all()
    checkin_rows = (await db.execute(
        select(CheckIn.guest_id).where(CheckIn.event_id == event_id)
    )).all()
    checked_in = {r[0] for r in checkin_rows}
    purchases = (await db.execute(
        select(TicketPurchase).where(
            TicketPurchase.event_id == event_id, TicketPurchase.status == "completed"
        ).order_by(TicketPurchase.created_at.desc())
    )).scalars().all()

    wb = Workbook()

    def setup_sheet(ws, title, headers, rows, widths=None):
        ws.title = title
        bold = Font(bold=True)
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = bold
            if widths and col <= len(widths):
                ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = widths[col - 1]
        for r, row in enumerate(rows, 2):
            for col, val in enumerate(row, 1):
                ws.cell(row=r, column=col, value=val)

    accepted = [g for g in guests if g.rsvp_status == "accepted"]
    summary_rows = [
        ["Event", event.title],
        ["Date", str(event.event_date)],
        ["Time", str(event.event_time)],
        ["Venue", event.venue or ""],
        ["Status", event.status],
        ["Visibility", "Public" if event.is_public else "Private"],
        ["Guest cap", event.guest_count_range],
        ["Registration", "Open" if event.registration_open else "Closed"],
        ["RSVP links", "Active" if getattr(event, "rsvp_open", True) else "Paused"],
        ["", ""],
        ["Total guests", len(guests)],
        ["Accepted", len(accepted)],
        ["Pending", len([g for g in guests if g.rsvp_status == "pending"])],
        ["Declined", len([g for g in guests if g.rsvp_status == "declined"])],
        ["Invites sent", len([g for g in guests if g.invite_sent])],
        ["Checked in", len(checked_in)],
        ["Ticket orders", len(purchases)],
        ["Tickets sold", sum(p.quantity or 0 for p in purchases)],
        ["Ticket revenue", sum(p.amount or 0 for p in purchases)],
    ]
    ws = wb.active
    setup_sheet(ws, "Summary", ["Metric", "Value"], summary_rows, [22, 50])

    guest_rows = []
    for g in guests:
        cd = g.custom_data or {}
        guest_rows.append([
            g.id, g.name, g.email or "", g.phone or "",
            g.category or "", cd.get("organization", ""),
            g.rsvp_status or "", "Yes" if g.invite_sent else "No",
            "Yes" if g.id in checked_in else "No",
            cd.get("reference", ""), g.notes or "",
        ])
    setup_sheet(wb.create_sheet(), "Guests",
                ["ID", "Name", "Email", "Phone", "Category", "Organization",
                 "RSVP", "Invited", "Checked in", "Payment ref", "Notes"],
                guest_rows, [8, 28, 32, 16, 18, 24, 10, 10, 12, 22, 30])

    sale_rows = [[p.buyer_name, p.buyer_email, p.buyer_phone or "",
                  p.package_name or "", p.quantity or 0, p.amount or 0,
                  p.reference, str(p.paid_at or p.created_at or "")]
                 for p in purchases]
    setup_sheet(wb.create_sheet(), "Ticket Sales",
                ["Buyer", "Email", "Phone", "Package", "Qty", "Amount", "Reference", "Paid at"],
                sale_rows, [28, 32, 16, 20, 8, 12, 24, 22])

    cats: dict[str, dict] = {}
    for g in guests:
        c = g.category or "Uncategorized"
        d = cats.setdefault(c, {"total": 0, "checked_in": 0})
        d["total"] += 1
        if g.id in checked_in:
            d["checked_in"] += 1
    setup_sheet(wb.create_sheet(), "Accreditation",
                ["Category", "Total", "Checked in"],
                [[c, d["total"], d["checked_in"]] for c, d in sorted(cats.items())],
                [28, 12, 14])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await log_action(db=db, user_id=admin.id, action="event_report_downloaded",
                     resource_type="event", resource_id=event_id,
                     description=f"Downloaded post-event report ({len(guests)} guests)")
    safe_title = "".join(ch if ch.isalnum() or ch in ("-", "_") else "_" for ch in (event.title or "event"))[:40]
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="event-{event_id}-{safe_title}-report.xlsx"'},
    )


class MarkPaidPublishRequest(BaseModel):
    make_public: bool = True


@router.post("/admin/events/{event_id}/mark-paid-publish")
async def admin_mark_paid_publish(
    event_id: int,
    req: MarkPaidPublishRequest,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Offline-payment rescue: complete pending non-resend payments and
    publish the event, mirroring what the Paystack webhook does."""
    from datetime import datetime, timezone as _tz
    from app.models.payment import Payment
    from app.services.notify import notify_subscribers

    event = await _admin_get_event(db, event_id)
    result = await db.execute(
        select(Payment).where(
            Payment.event_id == event_id,
            Payment.status == "pending",
            Payment.payment_type != "resend",
        )
    )
    pending = result.scalars().all()
    for p in pending:
        p.status = "completed"
        p.paid_at = datetime.now(_tz.utc)
    event.status = "published"
    if req.make_public:
        event.is_public = True
    await db.commit()
    try:
        await notify_subscribers(db, event.id)
    except Exception:
        pass
    await log_action(db=db, user_id=admin.id, action="event_marked_paid",
                     resource_type="event", resource_id=event_id,
                     description=f"Completed {len(pending)} pending payment(s), published (public={req.make_public})")
    return {"completed_payments": len(pending), "status": event.status, "is_public": event.is_public}


class AdminInviteMessageRequest(BaseModel):
    invite_subject: str | None = None
    invite_body: str | None = None


@router.put("/admin/events/{event_id}/invite-message")
async def admin_update_invite_message(
    event_id: int,
    req: AdminInviteMessageRequest,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Admin edit of the per-event invite subject/body guests receive."""
    event = await _admin_get_event(db, event_id)
    if req.invite_subject is not None:
        event.invite_subject = req.invite_subject
    if req.invite_body is not None:
        event.invite_body = req.invite_body
    await db.commit()
    await log_action(db=db, user_id=admin.id, action="invite_message_updated",
                     resource_type="event", resource_id=event_id,
                     description="Invite subject/body updated by admin")
    return {"invite_subject": event.invite_subject, "invite_body": event.invite_body}


class SuspendUserRequest(BaseModel):
    suspended: bool


@router.post("/admin/users/{user_id}/suspend")
async def admin_suspend_user(
    user_id: int,
    req: SuspendUserRequest,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Suspend / unsuspend an account. Suspended users cannot log in or call the API."""
    target = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    if target.role in ("admin", "super_admin"):
        raise HTTPException(status_code=403, detail="Cannot suspend another admin")
    if target.id == admin.id:
        raise HTTPException(status_code=403, detail="Cannot suspend yourself")
    target.is_active = not req.suspended
    await db.commit()
    await log_action(db=db, user_id=admin.id, action="user_suspended" if req.suspended else "user_unsuspended",
                     resource_type="user", resource_id=user_id,
                     description=f"{target.email} {'suspended' if req.suspended else 'unsuspended'} by admin")
    return {"id": target.id, "email": target.email, "is_active": target.is_active}


@router.post("/admin/users/{user_id}/impersonate")
async def admin_impersonate_user(
    user_id: int,
    response: Response,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Support login: swap the admin's session for the organizer's dashboard.

    Audited. Cannot target other admins. The admin must log back in
    afterwards to restore their own session.
    """
    from app.core.security import create_access_token
    from app.api.auth import set_auth_cookie

    target = (await db.execute(select(User).where(User.id == user_id))).scalar_one_or_none()
    if not target:
        raise HTTPException(status_code=404, detail="User not found")
    if target.role in ("admin", "super_admin"):
        raise HTTPException(status_code=403, detail="Cannot impersonate another admin")
    token = create_access_token({"sub": str(target.id)})
    set_auth_cookie(response, token)
    await log_action(db=db, user_id=admin.id, action="admin_impersonate",
                     resource_type="user", resource_id=user_id,
                     description=f"Admin {admin.email} opened organizer dashboard of {target.email}")
    return {"user": {"id": target.id, "email": target.email, "full_name": target.full_name},
            "message": f"Now viewing as {target.email}. Log back in to restore admin session."}


@router.get("/admin/events/{event_id}/review-details")
async def get_event_review_details(
    event_id: int,
    admin: User = Depends(check_admin),
    db: AsyncSession = Depends(get_db),
):
    """Get detailed review information for an event"""
    result = await db.execute(select(Event).where(Event.id == event_id))
    event = result.scalar_one_or_none()

    if not event:
        raise HTTPException(status_code=404, detail="Event not found")

    organizer_result = await db.execute(
        select(User).where(User.id == event.organizer_id)
    )
    organizer = organizer_result.scalar_one_or_none()

    return {
        "id": event.id,
        "title": event.title,
        "host_name": event.host_name,
        "event_type": event.event_type,
        "description": event.description,
        "event_date": event.event_date,
        "event_time": event.event_time,
        "venue": event.venue,
        "guest_count_range": event.guest_count_range,
        "organizer": {
            "id": organizer.id,
            "name": organizer.full_name,
            "email": organizer.email,
        } if organizer else None,
        "review_status": event.review_status,
        "flagged_keywords": event.flagged_keywords or [],
        "review_note": event.review_note,
        "created_at": event.created_at,
    }
